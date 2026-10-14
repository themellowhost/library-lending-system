import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Ultra-Modern Apple Dark Theme Palette
    BG_DARK = RGBColor(9, 9, 11)           # Pure Dark Background (#09090b)
    CARD_BG = RGBColor(20, 20, 24)         # Elevated Card Surface (#141418)
    CARD_BORDER = RGBColor(39, 39, 46)     # Subtle Dark Border (#27272e)
    ACCENT_BLUE = RGBColor(0, 113, 227)    # Apple Electric Blue (#0071E3)
    ACCENT_INDIGO = RGBColor(139, 92, 246) # Purple Accent (#8B5CF6)
    ACCENT_CYAN = RGBColor(56, 189, 248)   # Sky Cyan (#38BDF8)
    TEXT_MAIN = RGBColor(250, 250, 250)    # Primary Light Text (#FAFAFA)
    TEXT_MUTED = RGBColor(161, 161, 170)   # Slate Muted Text (#A1A1AA)
    EMERALD = RGBColor(16, 185, 129)       # Emerald Green (#10B981)
    AMBER = RGBColor(245, 158, 11)         # Amber Orange (#F59E0B)

    slides_data = [
        # SLIDE 1: Title Slide (Dark)
        {
            "type": "title",
            "tag": "DBMS SEMESTER 5 • LAB CONTINUOUS ASSESSMENT (LCA)",
            "title": "Library Lending System (LiberFlow)",
            "subtitle": "Relational Modeling, Interactive E-Reader, Federated OAuth & Apple-Inspired Architecture",
            "footer": "Engineered with React 19 • Node.js/TypeScript • SQLite (WAL Mode) • Open Library API • Deployed to GitHub Pages"
        },
        # SLIDE 2: Problem Statement & Fundamental Separation
        {
            "type": "content_2col",
            "tag": "SYSTEM REQUIREMENTS & CORE DOMAIN",
            "title": "Problem Statement & The Fundamental Design Choice",
            "left_title": "The Core Domain Challenge",
            "left_bullets": [
                "A library does not lend abstract titles; it lends physical barcoded items placed on distinct shelves.",
                "Naive systems merge catalog entries and physical stock, leading to severe data anomalies and inventory loss.",
                "Circulation demands strict transactional ACID guarantees to eliminate concurrent double-booking.",
                "Students require self-service reservations and loan extensions with duty librarian approval oversight."
            ],
            "right_title": "Key Design Choice: Title vs. Copy Separation",
            "right_bullets": [
                "Book Entity (Metadata): Title, ISBN, Publisher, Category, Authors, Digital Reader/PDF URLs.",
                "Copy Entity (Physical Item): Asset Barcode, Shelf Location, Status ('available', 'on_loan', 'on_hold').",
                "Scalability: 1 title seamlessly scales to 10+ physical copies across library stacks with independent circulation.",
                "Digital Hybrid: Physical copies circulate via barcode transactions while digital editions stream via online e-readers."
            ]
        },
        # SLIDE 3: Tech Stack Selection & Architecture
        {
            "type": "content_3card",
            "tag": "ARCHITECTURE DECISION RECORD (ADR)",
            "title": "Technology Stack: Selection Rationale & Advantages",
            "cards": [
                {
                    "title": "Frontend: React 19 + Apple UI",
                    "subtitle": "Apple Dark Glass & Interactive E-Reader",
                    "points": [
                        "Apple Dark Aesthetics: Apple glass (.apple-glass), SF Pro fonts, and pill buttons.",
                        "Fast E-Reader: Direct embedded 2-page reader canvas with 86vh viewport.",
                        "Federated OAuth: One-click Sign In with Apple and Google with official brand vector SVGs.",
                        "Static Deployment Resilience: In-memory runtime engine enabling full functionality on GitHub Pages."
                    ]
                },
                {
                    "title": "Backend: Node.js + Express + TS",
                    "subtitle": "Type Safety & Open API Integration",
                    "points": [
                        "Full TypeScript Stack: Shared data interfaces between client, server, and schema models.",
                        "Open Library Integration: Live ISBN auto-lookup and open-access search for millions of public books.",
                        "Bcrypt & JWT Auth: Salted password hashing with signed authorization tokens.",
                        "Hold Reaper Worker: Automated 60s background daemon promoting waiting hold queues."
                    ]
                },
                {
                    "title": "DBMS: Relational Engine (WAL)",
                    "subtitle": "ACID Transactions & Partial Unique Indexes",
                    "points": [
                        "Foreign Keys Enforced (PRAGMA foreign_keys = ON): Guaranteed referential integrity.",
                        "Write-Ahead Logging (WAL): Concurrency support for concurrent reads and writes.",
                        "Partial Unique Indexes: True DBMS constraints (1 active loan per copy, 1 active hold per member).",
                        "Audit Trail: Immutable audit_logs table for all checkouts, returns, extensions, and fine settlements."
                    ]
                }
            ]
        },
        # SLIDE 4: High-Resolution Rendered ER Diagram (Dark)
        {
            "type": "er_diagram_image",
            "tag": "RELATIONAL DATABASE DESIGN • ENTITY RELATIONSHIPS",
            "title": "Relational Schema & Entity-Relationship Architecture",
            "image_path": "/Users/akhshatkampassi/Documents/GitHub/DIV-21---B.TECH-ECE-AI-ML----Black-Box/Study/Labs/DBMS/Sem 5 LCA/er_diagram_dark.png"
        },
        # SLIDE 5: Transactional Checkout Workflow
        {
            "type": "content_steps",
            "tag": "WORKFLOW IMPLEMENTATION • TRANSACTION 1",
            "title": "Journey 1: Transactional Checkout Workflow",
            "steps": [
                {
                    "num": "01",
                    "head": "Input & Validation",
                    "desc": "Librarian scans Member Card (MEM-XXXXX) and Item Barcode. Zod validates payload format."
                },
                {
                    "num": "02",
                    "head": "Account & Quota Check",
                    "desc": "Verifies member is 'active' with zero unpaid fines and active loans < max_loans (quota limit)."
                },
                {
                    "num": "03",
                    "head": "Copy Availability Check",
                    "desc": "Ensures copy status is 'available' or allocated to this specific borrower via reservation hold."
                },
                {
                    "num": "04",
                    "head": "Atomic Lock & Commit",
                    "desc": "Inserts record into loans, updates copy status to 'on_loan', fulfills holds in one atomic db.transaction()."
                }
            ]
        },
        # SLIDE 6: Return Workflow & Dynamic Overdue Calculation
        {
            "type": "content_steps",
            "tag": "WORKFLOW IMPLEMENTATION • TRANSACTION 2",
            "title": "Journey 2: Return Workflow & Hold Queue Promotion",
            "steps": [
                {
                    "num": "01",
                    "head": "Barcode Check-In",
                    "desc": "Librarian scans returned copy barcode. System locates active loan record where returned_at IS NULL."
                },
                {
                    "num": "02",
                    "head": "Dynamic Overdue Calculation",
                    "desc": "Calculates (returned_at - due_at). If positive, computes fine ($1.00/day) and persists fine_amount."
                },
                {
                    "num": "03",
                    "head": "Reservation Queue Check",
                    "desc": "Queries holds for next 'waiting' member ordered by placed_at ASC (First-Come, First-Served)."
                },
                {
                    "num": "04",
                    "head": "State Transition",
                    "desc": "If hold exists: copy status -> 'on_hold', hold -> 'ready_for_pickup' (3-day window). Else: copy -> 'available'."
                }
            ]
        },
        # SLIDE 7: Loan Extension Application & Librarian Approval Workflow
        {
            "type": "content_2col",
            "tag": "LENDING EXTENSION & APPROVAL GOVERNANCE",
            "title": "Loan Extension Application & Librarian Approval Governance",
            "left_title": "Student / Member Application Flow",
            "left_bullets": [
                "Self-Service Application: Members click 'Apply for Extension' in My Account rather than automated renewal.",
                "Academic Justification: Member selects requested window (+7, +14, +21 days) and submits reasoning (exams, thesis).",
                "Instant Dialog Dismissal: Modal auto-dismisses upon submission with reactive UI banner.",
                "Live Lifecycle Badging: Dynamic badge changes to 'Under Librarian Review', 'Extension Approved', or 'Extension Rejected' with librarian notes."
            ],
            "right_title": "Librarian Approval & Inventory Protection",
            "right_bullets": [
                "Librarian Control Desk: Dedicated 'Pending Extension Applications' review queue on librarian dashboard.",
                "Demand & Queue Checks: Librarian inspects borrower history and waiting hold counts before deciding.",
                "Instant Queue Update: Approving/rejecting immediately removes the card from the pending queue.",
                "Quota & Overdue Guard: Blocked if item is already overdue or if 2-extension maximum has been reached."
            ]
        },
        # SLIDE 8: Multi-User Role Architecture & View Differentiation
        {
            "type": "content_3card",
            "tag": "ROLE-BASED ACCESS CONTROL (RBAC)",
            "title": "User Role Architecture: Distinct Views for Every User",
            "cards": [
                {
                    "title": "Member View (Students)",
                    "subtitle": "Self-Service Portal & Reservations",
                    "points": [
                        "Dedicated Portal: 'My Borrowed & Holds' showing due dates, fine alerts, and queue rank.",
                        "Online Reservation: 'Reserve for Library Pickup' with 3-day hold window on available copies.",
                        "Extension Applications: Formal loan window prolongation request with live approval feedback.",
                        "Open E-Library: Browse open-access books and open full-screen e-readers."
                    ]
                },
                {
                    "title": "Librarian View (Staff)",
                    "subtitle": "Inventory Health & Circulation Desk",
                    "points": [
                        "5 Key Operational Metrics: Total Titles, Shelf Availability %, Active Loans, Ready for Pickup, Extension Requests.",
                        "Extension Approval Desk: One-click Approve (+Xd) and Reject interface with student justification view.",
                        "Circulation Desk: High-speed barcode scanning interface for atomic Checkouts and Returns.",
                        "ISBN Auto-Scanning: Scan or enter ISBN to auto-fetch catalog metadata from Open Library."
                    ]
                },
                {
                    "title": "Administrator View (IT/Admin)",
                    "subtitle": "System Governance & Audit Trail",
                    "points": [
                        "Member Roster: Provision student accounts, mint MEM-XXXXX cards, toggle active/suspended.",
                        "Circulation Analytics: Overdue delinquency reports and fine collection tracking ($1/day).",
                        "Security Oversight: Role-based authorization rules and password hash verification.",
                        "Audit Ledger: Complete historical trail of all checkout, return, extension, and fine events."
                    ]
                }
            ]
        },
        # SLIDE 9: Open-Access E-Reader & Global Archive Integration
        {
            "type": "content_2col",
            "tag": "DIGITAL EXTENSION • OPEN SOURCE INTEGRATION",
            "title": "Open-Access E-Reader & Global Archive Integration",
            "left_title": "Open Library & Internet Archive Integration",
            "left_bullets": [
                "Public Domain Search: Proxied GET /api/books/open-search queries Open Library's Search API for full-text items.",
                "Direct IA Resolution: Maps Internet Archive identifiers to direct embed streams (https://archive.org/embed/{ia_id}).",
                "Auto-Cataloging: Scanned ISBNs auto-resolve metadata (publishers, edition year, authors, cover images).",
                "Curated Classics: Includes full editions of SICP, Turing (1936), Algorithms, and Computer Organization."
            ],
            "right_title": "Custom Embedded Reader Experience",
            "right_bullets": [
                "Maximized Reading Viewport: Expanded 86vh iframe height and max-w-[1550px] layout.",
                "Embed URL Parameter: Appending ?ui=embed removes redundant headers for a fast, clean book spread.",
                "Multi-Modal Formats: Direct toggle between interactive 2-page e-reader canvas and downloadable PDF edition.",
                "Zero Auth Wall: Free open-access books accessible to guest visitors without mandatory login."
            ]
        },
        # SLIDE 10: Summary & Live Demonstration
        {
            "type": "content_2col",
            "tag": "DEPLOYMENT & PROJECT CONCLUSION",
            "title": "System Summary & Verification Details",
            "left_title": "DBMS Academic Requirements Met",
            "left_bullets": [
                "Normalization: Third Normal Form (3NF) relational design with zero data redundancy.",
                "ACID Concurrency: Write-Ahead Logging (WAL), atomic transactions, and partial unique indexes.",
                "Integrity Constraints: ON DELETE CASCADE, CHECK constraints, and non-nullable keys.",
                "Auditability: Immutable circulation ledger tracking every checkout, return, extension, and fine."
            ],
            "right_title": "Live Production Deployment",
            "right_bullets": [
                "Production App: https://themellowhost.github.io/library-lending-system/",
                "GitHub Repository: https://github.com/themellowhost/library-lending-system",
                "Full Persona Testing: One-click switcher between Member (Alex), Librarian (Elena), and Admin (Marcus).",
                "Dual-Mode Architecture: Full operational capability on static hosting via resilient client fallback engine."
            ]
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_slide_layout)
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK

        if data["type"] == "title":
            # Top accent stripe
            top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1.5), Inches(0.08))
            top_bar.fill.solid()
            top_bar.fill.fore_color.rgb = ACCENT_BLUE
            top_bar.line.fill.background()

            txBox = slide.shapes.add_textbox(Inches(1), Inches(1.4), Inches(11.3), Inches(4.5))
            tf = txBox.text_frame
            tf.word_wrap = True

            p_tag = tf.paragraphs[0]
            p_tag.text = data["tag"]
            p_tag.font.size = Pt(13)
            p_tag.font.bold = True
            p_tag.font.color.rgb = ACCENT_BLUE
            p_tag.space_after = Pt(20)

            p_title = tf.add_paragraph()
            p_title.text = data["title"]
            p_title.font.size = Pt(44)
            p_title.font.bold = True
            p_title.font.color.rgb = TEXT_MAIN
            p_title.space_after = Pt(16)

            p_sub = tf.add_paragraph()
            p_sub.text = data["subtitle"]
            p_sub.font.size = Pt(20)
            p_sub.font.color.rgb = ACCENT_CYAN
            p_sub.space_after = Pt(36)

            p_foot = tf.add_paragraph()
            p_foot.text = data["footer"]
            p_foot.font.size = Pt(12)
            p_foot.font.color.rgb = TEXT_MUTED

        elif data["type"] == "er_diagram_image":
            add_header(slide, data["tag"], data["title"], ACCENT_BLUE, TEXT_MAIN)

            # Insert high-resolution rendered dark ER diagram image
            img_path = data["image_path"]
            if os.path.exists(img_path):
                slide.shapes.add_picture(
                    img_path,
                    Inches(0.8),
                    Inches(1.7),
                    width=Inches(11.73),
                    height=Inches(5.35)
                )

        elif data["type"] == "content_2col":
            add_header(slide, data["tag"], data["title"], ACCENT_BLUE, TEXT_MAIN)

            # Left Card
            add_card_with_bullets(
                slide,
                x=Inches(0.8),
                y=Inches(1.8),
                w=Inches(5.7),
                h=Inches(5.1),
                card_bg=CARD_BG,
                card_border=CARD_BORDER,
                header_text=data["left_title"],
                bullets=data["left_bullets"],
                accent_color=ACCENT_BLUE,
                text_color=TEXT_MAIN,
                muted_color=TEXT_MUTED
            )

            # Right Card
            add_card_with_bullets(
                slide,
                x=Inches(6.8),
                y=Inches(1.8),
                w=Inches(5.7),
                h=Inches(5.1),
                card_bg=CARD_BG,
                card_border=CARD_BORDER,
                header_text=data["right_title"],
                bullets=data["right_bullets"],
                accent_color=ACCENT_INDIGO,
                text_color=TEXT_MAIN,
                muted_color=TEXT_MUTED
            )

        elif data["type"] == "content_3card":
            add_header(slide, data["tag"], data["title"], ACCENT_BLUE, TEXT_MAIN)

            card_w = Inches(3.75)
            gap = Inches(0.25)
            start_x = Inches(0.8)

            for i, card_info in enumerate(data["cards"]):
                cx = start_x + (i * (card_w + gap))
                cy = Inches(1.8)
                ch = Inches(5.1)

                add_card_with_points(
                    slide,
                    x=cx,
                    y=cy,
                    w=card_w,
                    h=ch,
                    card_bg=CARD_BG,
                    card_border=CARD_BORDER,
                    title=card_info["title"],
                    subtitle=card_info["subtitle"],
                    points=card_info["points"],
                    accent_color=ACCENT_BLUE if i == 0 else (ACCENT_INDIGO if i == 1 else EMERALD),
                    text_color=TEXT_MAIN,
                    muted_color=TEXT_MUTED
                )

        elif data["type"] == "content_steps":
            add_header(slide, data["tag"], data["title"], ACCENT_BLUE, TEXT_MAIN)

            step_w = Inches(2.78)
            gap = Inches(0.2)
            start_x = Inches(0.8)

            for i, step_info in enumerate(data["steps"]):
                sx = start_x + (i * (step_w + gap))
                sy = Inches(1.8)
                sh = Inches(5.1)

                add_step_card(
                    slide,
                    x=sx,
                    y=sy,
                    w=step_w,
                    h=sh,
                    card_bg=CARD_BG,
                    card_border=CARD_BORDER,
                    step_num=step_info["num"],
                    head=step_info["head"],
                    desc=step_info["desc"],
                    accent_color=ACCENT_BLUE if i % 2 == 0 else ACCENT_CYAN,
                    text_color=TEXT_MAIN,
                    muted_color=TEXT_MUTED
                )

    output_path = "/Users/akhshatkampassi/Documents/GitHub/DIV-21---B.TECH-ECE-AI-ML----Black-Box/Study/Labs/DBMS/Sem 5 LCA/Library_Lending_System_Presentation.pptx"
    prs.save(output_path)
    print(f"Dark-theme presentation with ER diagram saved successfully to: {output_path}")

def add_header(slide, tag_text, title_text, accent_color, text_color):
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.4))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = accent_color

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(26)
    p_title.font.bold = True
    p_title.font.color.rgb = text_color

def add_card_with_bullets(slide, x, y, w, h, card_bg, card_border, header_text, bullets, accent_color, text_color, muted_color):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = card_bg
    card.line.color.rgb = card_border
    card.line.width = Pt(1.2)

    top_stripe = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.08))
    top_stripe.fill.solid()
    top_stripe.fill.fore_color.rgb = accent_color
    top_stripe.line.fill.background()

    tb = slide.shapes.add_textbox(x + Inches(0.35), y + Inches(0.3), w - Inches(0.7), h - Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p_head = tf.paragraphs[0]
    p_head.text = header_text
    p_head.font.size = Pt(17)
    p_head.font.bold = True
    p_head.font.color.rgb = text_color
    p_head.space_after = Pt(16)

    for bullet in bullets:
        p = tf.add_paragraph()
        p.text = "• " + bullet
        p.font.size = Pt(12)
        p.font.color.rgb = muted_color
        p.space_after = Pt(11)

def add_card_with_points(slide, x, y, w, h, card_bg, card_border, title, subtitle, points, accent_color, text_color, muted_color):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = card_bg
    card.line.color.rgb = card_border
    card.line.width = Pt(1.2)

    top_stripe = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.08))
    top_stripe.fill.solid()
    top_stripe.fill.fore_color.rgb = accent_color
    top_stripe.line.fill.background()

    tb = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(0.25), w - Inches(0.6), h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(15)
    p_title.font.bold = True
    p_title.font.color.rgb = text_color
    p_title.space_after = Pt(2)

    p_sub = tf.add_paragraph()
    p_sub.text = subtitle
    p_sub.font.size = Pt(10)
    p_sub.font.bold = True
    p_sub.font.color.rgb = accent_color
    p_sub.space_after = Pt(14)

    for pt in points:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = muted_color
        p.space_after = Pt(9)

def add_step_card(slide, x, y, w, h, card_bg, card_border, step_num, head, desc, accent_color, text_color, muted_color):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = card_bg
    card.line.color.rgb = card_border
    card.line.width = Pt(1.2)

    tb = slide.shapes.add_textbox(x + Inches(0.25), y + Inches(0.3), w - Inches(0.5), h - Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p_num = tf.paragraphs[0]
    p_num.text = step_num
    p_num.font.size = Pt(30)
    p_num.font.bold = True
    p_num.font.color.rgb = accent_color
    p_num.space_after = Pt(12)

    p_head = tf.add_paragraph()
    p_head.text = head
    p_head.font.size = Pt(15)
    p_head.font.bold = True
    p_head.font.color.rgb = text_color
    p_head.space_after = Pt(12)

    p_desc = tf.add_paragraph()
    p_desc.text = desc
    p_desc.font.size = Pt(11.5)
    p_desc.font.color.rgb = muted_color
    p_desc.line_spacing = 1.2

if __name__ == "__main__":
    create_deck()
